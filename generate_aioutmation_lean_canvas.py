"""Generate Lean Canvas Word doc for AIOutmation product."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_bg(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def style_run(run, size=10, bold=False, color=None, font_name="Yu Gothic"):
    run.font.size = Pt(size)
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def write_cell(cell, text, size=9, bold=False, bg=None, color=None, align=None):
    cell.text = ""
    if bg:
        set_cell_bg(cell, bg)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        if align:
            p.alignment = align
        run = p.add_run(line)
        style_run(run, size=size, bold=bold, color=color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level <= 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    sizes = {0: 22, 1: 18, 2: 14, 3: 12}
    colors = {0: "1F3A5F", 1: "1F3A5F", 2: "2E5C8A", 3: "444444"}
    style_run(run, size=sizes.get(level, 11), bold=True, color=colors.get(level, "000000"))


def add_paragraph(doc, text, size=10, bold=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, color=color)


def add_callout(doc, title, body, color="FFF4E0"):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_bg(cell, color)
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(3)
    r = p1.add_run(title)
    style_run(r, size=11, bold=True, color="C04000")
    p2 = cell.add_paragraph()
    r = p2.add_run(body)
    style_run(r, size=10)


def setup_doc(doc):
    section = doc.sections[0]
    new_w, new_h = section.page_height, section.page_width
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = new_w
    section.page_height = new_h
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    style = doc.styles["Normal"]
    style.font.name = "Yu Gothic"
    style.font.size = Pt(10)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Yu Gothic")


def cover_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("AIOutmation")
    style_run(r, size=48, bold=True, color="1F3A5F")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("リーンキャンバス")
    style_run(r, size=24, bold=True, color="2E5C8A")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    r = p.add_run("AIO最適化を、書く前から書いた後まで、自動で回す。")
    style_run(r, size=14, color="555555")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    r = p.add_run("──── アウトプット② ────")
    style_run(r, size=11, color="888888")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("ペイン分析 + Chapter 4 PDCA自動化分析 + リーンキャンバス本体")
    style_run(r, size=11, color="555555")
    doc.add_page_break()


# ====== Lean Canvas Data ======

LC = {
    "problem": """1) AI検索可視性の測定が手作業地獄(月12.5分×多モデル×多プロンプト=実質数時間)
2) 計測→改修(コンテンツ修正・Schema実装・CMS入稿)が完全分断、4部門横断で1施策2-6週間
3) AIO業界SOPが日本市場に存在せず、各社が独自にゼロから設計→属人化・知識流出
4) 経営/クライアントへのROI証明が困難(Visibility→売上の因果が示せない、業界ベンチマーク無し)
5) ノイズ(LLM揺らぎ)とシグナルの分離が統計的にできず、誤検知or見逃しが頻発
6) 多言語・多リージョン・サブブランドへのスケール対応で人手が爆発(代理店は3社目で破綻)
7) 既存ベンチマークツール(Profound/Peec/Passionfruit)は計測特化で改修・PR・経営報告まで届かない""",

    "customer_segments": """【主要顧客①: 上場会社のWebマーケティング部門】
- 想定: 東証プライム/グロース市場上場の事業会社、ARR/売上100億円〜数千億円規模
- 業種優先順位: B2B SaaS > 金融・保険 > 大手リテール・EC > 製薬・ヘルスケア > 通信・メディア
- 購買決裁ライン: マーケ部長/Director → CMO/VP Marketing → CFO/取締役会
- 関与チーム: SEO/コンテンツ/PR/Web開発/法務/情シス/購買
- 個人ペルソナA: Sr. Manager(GEO/AEO Lead)、30代後半、KGIはパイプライン貢献+ブランドVisibility
- 既存ツール: GA4 / Search Console / Semrush / HubSpot / Salesforce / Looker

【主要顧客②: ネット広告代理店(上場企業をクライアントとする)】
- 想定: 50-300名規模、SEO/コンテンツ/運用広告を展開、月額リテイナー型
- 顧客層: 上場B2B SaaS 5-12社/担当、月額$3k-$30k(40-450万円)
- 購買決裁ライン: Group Director / Partner
- 個人ペルソナB: Account Director/GEO Lead、KGIはリテンション+新規受注+リテイナー単価
- 既存ツール: Semrush / Ahrefs / GA4 / Looker Studio / AgencyAnalytics / Asana

【副次セグメント(将来拡張)】: 国内Tier 2のコンサルティングファーム、PR代理店、SI企業のマーケDX部門""",

    "uvp": """瀧内本「AIO教科書」が示す体系を、ソフトウェアで完全実装。
Webマーケ・コンテンツ・SEO・PRの分断作業をワンプラットフォームで統合し、
測定→改修→効果検証→経営報告までを自動化する、日本市場初の統合型AIO基盤。

※高位概念: 「年間1,500時間の手作業を100時間に圧縮し、AI時代のマーケROIを再定義する」""",

    "high_level_concept": "Webマーケのための『AIO版GitHub Copilot+Datadog』。書く前(Plan)から書いた後(運用)まで、AIOピラミッドの全層を自動化。",

    "solution": """主要機能8本(AIOピラミッドの全層をカバー):

①【SEO土台】SEO×AIO監査エンジン
   - 既存サイトのHTML階層・構造化データ・E-E-A-T充足度を自動診断
   - 90日ロードマップ自動生成

②【AEO理解層】文章構成最適化AI(Ch.2準拠)
   - PREP/FAQ/HowTo構文への自動リライト
   - 既存ページを取り込みAEO対応版を提案

③【AEO/GEO】プロンプト自動生成エンジン(Ch.1+4準拠)
   - 自社URL+CRM(Salesforce/HubSpot)+営業会話(Gong/Zoom)から
     ICP×購買ファネル×製品でMoney Prompt 30-1,000本を自動生成
   - 「なぜこのプロンプトか」をAIが自動文書化(属人化解消)

④【GEO引用層】文書構造化エンジン(Ch.3準拠)
   - JSON-LD自動生成(FAQPage/HowTo/Article/Person/Organization)
   - Evidence+Reference自動挿入(citeタグ整備)
   - GitHub/Webflow/WordPress API連携でCMSにPR自動起票

⑤【運用】引用率測定エンジン(Ch.4 PDCA完全自動化)
   - LLM API+UIスクレイピングで毎日プロンプト実行
   - Citation Rate自動算出、FAQ管理表/記録表/集計表の自動更新
   - 70%閾値を割ったら自動アラート、原因分析AIが改修候補提示
   - ★Signal/Noise Engine: LLM揺らぎを統計的に除去(複数回実行+信頼区間)

⑥【LLMO学習層】知識資産化アシスタント(Ch.5準拠)
   - llms.txt自動生成・更新
   - Perplexity/Gemini/ChatGPT別の最適化レコメンド
   - 更新頻度・一貫性スコアを継続監視

⑦【経営報告】Executive Reporter
   - CMO向け3行サマリ自動生成(前月比/カテゴリ/競合差分)
   - QBR用デック自動作成、ROI物語AI(過去成功事例DBから類似ストーリー提示)
   - 業界ベンチマーク同梱(自社蓄積データ+提携データから業種別中央値提示)

⑧【代理店向け】マルチテナント運用
   - クライアント別ホワイトラベル(ロゴ/カラー/独自ドメイン/送信元メール)
   - クレジット制(プロンプト枠を流動配分)、全プラン無制限シート
   - Pitch環境(無料即発行、契約条件で自動失効)""",

    "channels": """【主要チャネル】
1. ダイレクトセールス(エンタープライズ向け): 上場企業マーケ部門への営業/PoC/CSM
2. パートナーシップ販売: 大手ネット広告代理店との戦略提携(売上シェアモデル)
3. 著者・専門家ネットワーク: AIO教科書著者瀧内氏/SEO業界のオピニオンリーダーとの連携
4. コンテンツマーケ: ベンチマークレポート・業界調査・Webinar公開
5. PLG(セルフサーブ): 下位プランで個人事業主・SMBを獲得→アップセル

【サブチャネル】
6. 業界カンファレンス(MarkeZine Day, AdTech Tokyo, Adobe Summit Japan等)
7. プロダクト統合パートナー(HubSpot Marketplace, Salesforce AppExchange)
8. メディア掲載(日経クロストレンド, MarkeZine, Web担当者Forum)""",

    "revenue_streams": """【SaaSサブスクリプション(主軸)】
- スターター: 月額3万円〜(50プロンプト、3LLMモデル、SMB向けセルフサーブ)
- グロース: 月額10万円(150プロンプト、5LLM、CMS連携1個)
- ビジネス: 月額30万円(500プロンプト、全LLM、CMS連携無制限、CSM)
- エンタープライズ: 月額50-200万円(無制限プロンプト、SOC 2、SSO、専任CSM、API)
- 代理店プラン: 月額20万円〜(クレジット制、ホワイトラベル、Pitch環境、無制限シート)

【補助収益】
- オンボーディング/コンサルティング: スポット50-300万円
- カスタムCMSコネクタ開発: スポット100-500万円
- 業界ベンチマークレポート販売: 単発10-50万円

【KPI想定】: 24ヶ月で ARR 5億円(エンプラ50社 + 代理店30社 + SMB 200社)""",

    "cost_structure": """【固定費】
- 開発人件費: エンジニア5-10名 (年4,000-8,000万円)
- セールス・CSM人件費: 5-10名 (年4,000-8,000万円)
- 経営・管理: 法務・経理・HR (年2,000万円)
- オフィス・SaaS等管理費: (年500-1,000万円)

【変動費】
- LLM API利用料: 1顧客あたり月1-5万円(プロンプト数依存)
- UIスクレイピング基盤運用: 月100-500万円(プロキシ・ブラウザ自動化)
- データセンター/クラウド: 月50-300万円(成長依存)

【一時費用】
- SOC 2 Type II取得: 1,500-3,000万円(初年度)
- 業界ベンチマーク用データ収集パートナーシップ: 500-1,000万円
- PR/マーケティング初期投資: 1,000-3,000万円

【損益分岐点想定】: ARR 3億円(エンプラ30社相当)""",

    "key_metrics": """【グロース指標】
- ARR (Annual Recurring Revenue): 24ヶ月で5億円目標
- 上場企業ロゴ獲得数 (Tier 1顧客)
- 代理店パートナー数 / 管理クライアント数
- NRR (Net Revenue Retention): 120%以上

【プロダクト指標】
- 顧客平均Citation Rate改善量(導入前→導入後3ヶ月)
- 月間プロンプト実行数(プラットフォーム全体)
- 自動生成されたCMS PR数 / 採用率
- 経営報告自動生成回数 / 顧客満足度

【効率指標】
- CAC回収期間 (Payback Period): 12ヶ月以内
- LTV/CAC比率: 3倍以上
- セルフサーブ→有料転換率: 5%以上
- エンタープライズ商談勝率: 30%以上""",

    "unfair_advantage": """【1】日本市場の体系的SOP実装で先行(参入障壁)
   - 瀧内賢氏「AIO教科書」(2026年3月技術評論社, SEO書籍7作+AI書籍8作の実績)を
     プロダクトとして実装する世界初のソフトウェア
   - 著者監修・推奨・共同マーケティングが成立すれば差別化決定打

【2】日本語LLM完全対応(Profound/Peec/Passionfruitの最弱点を突く)
   - Felo, Genspark, PKSHA, 国産LLMの完全カバー
   - Perplexity日本語/ChatGPT日本語/Gemini日本語の最適化分離

【3】Signal/Noise Engine(統計的処理)
   - LLM評価の研究フロンティア(Ai2 framework等)を商用実装
   - 競合3社は閾値ベース、当社は信頼区間付き判定で誤検知激減

【4】計測→改修PR→効果検証のクローズドループ
   - 競合3社は計測特化、当社のみCMS PR自動起票・改修ロジック蓄積まで
   - 「測って終わり」を脱する唯一のプラットフォーム

【5】業界ベンチマーク独自蓄積データ
   - 顧客横断データ(匿名化)から業種別中央値・上位10%を提示
   - 経営説得力の決定的な武器

【6】上場企業×代理店の同時最適化
   - 一般的なSaaSは「直販 vs 代理店」を二者択一にするが、
     当社は同一プラットフォームで両セグメントに最適化(クレジット制+ホワイトラベル+Pitch)""",
}


# ====== Build Document ======

doc = Document()
setup_doc(doc)
cover_page(doc)

# === Section 1: Pain Mapping ===
add_heading(doc, "ステップ① ペイン × AIO教科書施策 マッピング", level=1)
add_paragraph(doc, "前回分析した8つのペインを、PDF施策で機能化したときの解決度を評価。", size=11)

t = doc.add_table(rows=9, cols=4)
t.autofit = False
widths = [Cm(5.5), Cm(8.0), Cm(11.0), Cm(3.5)]
for r in t.rows:
    for i, c in enumerate(r.cells):
        c.width = widths[i]

headers = ["ペイン", "PDFの該当施策", "機能化案", "解決度"]
for i, h in enumerate(headers):
    write_cell(t.cell(0, i), h, size=11, bold=True, bg="1F3A5F", color="FFFFFF")

mapping = [
    ("プロンプト初期設計", "Ch.1 Query→Intent→Answer\nCh.4 Money Prompt選定",
     "ICP×競合×営業会話を入力としたAIプロンプト自動生成。Ch.1のQuery分類体系をテンプレ化。",
     "◎ 完全解決"),
    ("コンテンツ改修工数", "Ch.2 PREP→FAQ→HowTo\nCh.3 JSON-LD・Schema",
     "既存ページを取り込み、PREP/FAQ/HowTo構造へ自動リライト+JSON-LD自動生成。",
     "◎ 完全解決"),
    ("レポートコメント執筆", "Ch.4 Citation Rateトレンド分析",
     "引用率推移+競合差分+施策履歴をテンプレに流し込みAI所感生成。",
     "○ 大部分解決"),
    ("戦略立案・解釈", "Ch.1 AIOピラミッド\n90日ロードマップ",
     "現状診断 → ピラミッド層別ギャップ可視化 → 次の打ち手レコメンド。",
     "○ 大部分解決"),
    ("「なぜこのプロンプトか」文脈記録", "Ch.4 FAQ管理表「改善メモ」",
     "プロンプト追加時に「対象ICP/ファネル/根拠データ」をAIが自動文書化、変更履歴も保持。",
     "◎ 完全解決"),
    ("改修ロジック・勝ちパターン共有", "Ch.4 PDCA Act段階\nCh.5 Evidence深化",
     "引用率改善した施策を業界横断で蓄積→「勝ちパターンライブラリ」として配信。",
     "◎ 完全解決"),
    ("計測→CMS改修PRの自動起票", "Ch.3 HTML階層・dlタグ・Schema仕様",
     "GitHub/Headless CMS API連携で改修内容を自動PR起票(Schema・FAQ・原子段落)。",
     "◎ 完全解決"),
    ("コンテンツ・改修・PR・SEO統合", "Ch.1 AIOピラミッド全体\nCh.4 運用",
     "AIOピラミッドの5層(SEO+AEO+GEO+LLMO+運用)を1プラットフォームで完結。",
     "◎ 完全解決"),
]

for i, (pain, pdf, feat, deg) in enumerate(mapping):
    write_cell(t.cell(i+1, 0), pain, size=10, bold=True, bg="E8EEF7")
    write_cell(t.cell(i+1, 1), pdf, size=9)
    write_cell(t.cell(i+1, 2), feat, size=9, bg="F0F8E8")
    write_cell(t.cell(i+1, 3), deg, size=11, bold=True, color="2E7A4E", align=WD_ALIGN_PARAGRAPH.CENTER)

add_paragraph(doc, "→ 8ペインのうち6つは完全解決、2つは大部分解決。これは強いPMF仮説。",
              size=11, bold=True, color="2E7A4E")

doc.add_page_break()

# === Section 2: PDCA Automation Analysis ===
add_heading(doc, "ステップ② Chapter 4 PDCAサイクル 自動化分析", level=1)
add_callout(doc, "結論", "完全に半自動化可能。むしろ自動化されるべき設計。書籍はGoogle Sheets手動運用を前提とするが、全工程がソフトウェアで代替可能。",
           color="E0F4E0")

add_heading(doc, "PDCA各段階の自動化マッピング", level=2)
t = doc.add_table(rows=5, cols=4)
t.autofit = False
widths = [Cm(2.5), Cm(8.5), Cm(13.0), Cm(4.0)]
for r in t.rows:
    for i, c in enumerate(r.cells):
        c.width = widths[i]
headers = ["段階", "書籍の手動手順", "ソフトウェア半自動化案", "自動化度"]
for i, h in enumerate(headers):
    write_cell(t.cell(0, i), h, size=11, bold=True, bg="1F3A5F", color="FFFFFF")

pdca = [
    ("Plan", "FAQ管理表をExcelで作成、優先順位を手動整理",
     "プロンプトDB+自動優先順位スコアリング(前月引用率×ROIインパクト)", "90% 自動"),
    ("Do", "ライターにブリーフ手書き、Web開発に依頼、CMSに手動入稿",
     "AIブリーフ自動生成、CMS PR自動起票、ライター承認ワークフロー", "70% 自動"),
    ("Check", "Perplexity/AI Overviews/ChatGPTに月25回手動入力、結果をSheets手書き",
     "LLM API+UIスクレイピング+自動引用判定エンジン+データ自動投入", "100% 自動"),
    ("Act", "引用率を見て「どこを改修するか」を人間が考える",
     "引用率低下要因をAIが自動分析、改修候補を優先度付きで提示", "80% 自動"),
]
for i, (stage, manual, auto, deg) in enumerate(pdca):
    write_cell(t.cell(i+1, 0), stage, size=12, bold=True, bg="E8EEF7", align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(t.cell(i+1, 1), manual, size=9)
    write_cell(t.cell(i+1, 2), auto, size=9, bg="F0F8E8")
    write_cell(t.cell(i+1, 3), deg, size=11, bold=True, color="2E7A4E", align=WD_ALIGN_PARAGRAPH.CENTER)

add_heading(doc, "Chapter 4運用工数 Before/After", level=2)
t = doc.add_table(rows=7, cols=4)
t.autofit = False
widths = [Cm(7.0), Cm(5.5), Cm(5.5), Cm(4.0)]
for r in t.rows:
    for i, c in enumerate(r.cells):
        c.width = widths[i]
headers = ["工程", "手動(書籍通り)", "半自動化後", "削減率"]
for i, h in enumerate(headers):
    write_cell(t.cell(0, i), h, size=11, bold=True, bg="1F3A5F", color="FFFFFF")
hours = [
    ("月次AI引用確認", "12.5分(25回×30秒)", "0分(バックグラウンド自動)", "100%"),
    ("FAQ記録表入力", "15分", "0分", "100%"),
    ("FAQ集計表作成", "30分(月末)", "0分", "100%"),
    ("ヒートマップ作成", "60分", "0分", "100%"),
    ("原因分析・改善計画", "120分", "30分(AI候補を人がレビュー)", "75%"),
    ("【月間合計】", "約4時間", "約30分", "87.5%削減"),
]
for i, row in enumerate(hours):
    is_total = "合計" in row[0]
    bg = "FFF4E0" if is_total else None
    for ci, val in enumerate(row):
        write_cell(t.cell(i+1, ci), val, size=10, bold=is_total, bg=bg or ("E8EEF7" if ci == 0 else None),
                   align=WD_ALIGN_PARAGRAPH.CENTER if ci > 0 else None)

add_paragraph(doc, "→ 10クライアント並行時、代理店は月40時間→5時間。これだけでROIが出る。",
              size=11, bold=True, color="2E7A4E")

doc.add_page_break()

# === Section 3: Lean Canvas ===
add_heading(doc, "ステップ③ AIOutmation リーンキャンバス", level=1)
add_paragraph(doc, "Ash Maurya標準フォーマット(9ブロック)。ターゲット顧客・ペイン・優位性・ソリューションを特に詳細記述。",
              size=11)

# Lean Canvas as a table (3x3 layout, with adjusted spans)
# Top row: Problem | Solution | UVP | Unfair Advantage | Customer Segments
# Bottom row: Key Metrics | Channels  (under Solution)  // Cost Structure | Revenue Streams (full width)

# Let me use a simpler structure: vertical stacked sections with prominent headers
# This is more readable for Word format

def add_lc_section(title, body, color, accent_color):
    add_heading(doc, title, level=2)
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_bg(cell, color)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(body)
    style_run(r, size=10)


# Customer Segments first (most important per user request)
add_lc_section("② 顧客セグメント (Customer Segments)", LC["customer_segments"], "FFF8E0", "C04000")

# Problem
add_lc_section("① 課題 (Problem) - Top 7", LC["problem"], "FFE0E0", "C04000")

# UVP
add_heading(doc, "③ 独自価値提案 (Unique Value Proposition)", level=2)
add_callout(doc, "UVP", LC["uvp"], color="E0F4E0")
add_paragraph(doc, f"ハイレベル概念: {LC['high_level_concept']}", size=10, bold=True, color="2E5C8A")

doc.add_page_break()

# Solution
add_lc_section("④ ソリューション (Solution) - 主要機能8本", LC["solution"], "F0F8FF", "1F3A5F")

doc.add_page_break()

# Unfair Advantage
add_lc_section("⑨ 圧倒的な優位性 (Unfair Advantage)", LC["unfair_advantage"], "FFF8E0", "C04000")

doc.add_page_break()

# Channels
add_lc_section("⑤ チャネル (Channels)", LC["channels"], "F5F5F5", "555555")

# Revenue Streams
add_lc_section("⑥ 収益の流れ (Revenue Streams)", LC["revenue_streams"], "F5F5F5", "555555")

doc.add_page_break()

# Cost Structure
add_lc_section("⑦ コスト構造 (Cost Structure)", LC["cost_structure"], "F5F5F5", "555555")

# Key Metrics
add_lc_section("⑧ 主要指標 (Key Metrics)", LC["key_metrics"], "F5F5F5", "555555")

doc.add_page_break()

# === Summary ===
add_heading(doc, "総括: AIOutmationの戦略ポジション", level=1)

add_heading(doc, "コアの差別化ストーリー", level=2)
add_paragraph(doc,
"""AIOutmationは、瀧内本「AIO教科書」が示す体系を、ソフトウェアで完全実装するプロダクトです。

書籍が前提とする「Google Sheets手動運用」を、データベース+LLM API+CMS連携によって自動化します。
これにより、月12.5分×複数モデル×多数プロンプト=実質数時間〜日の手作業を、月30分のAI提案レビューに圧縮します。

競合(Profound/Peec AI/Passionfruit Labs)は計測特化で「測って終わり」ですが、AIOutmationは
測定→改修PR→効果検証→経営報告までクローズドループで自動化する点で根本的に異なります。

加えて、日本市場で先行(英語ツール後追いではない)、日本語LLM完全対応、Signal/Noise Engineによる
統計的ノイズ除去、業界ベンチマーク独自蓄積、上場企業×代理店の同時最適化という5つの参入障壁を
備えています。""",
              size=11)

add_heading(doc, "次のアウトプット予告", level=2)
add_paragraph(doc,
"""アウトプット③ MVP機能仕様書(委託先ソフトウェア開発会社とのディスカッション資料)では、
本リーンキャンバスのソリューション8機能から、初回リリースに含めるべきMVP機能を3-5個に絞り、
- 各機能の目的・スコープ・成功基準
- 技術スタック想定(LLM API選定、UIスクレイピング基盤、CMS連携)
- データモデル概要
- リリース順序とマイルストーン
を提示します。""",
              size=11)

doc.save("/home/user/AIO/AIOutmation_LeanCanvas.docx")
print("Saved: AIOutmation_LeanCanvas.docx")
